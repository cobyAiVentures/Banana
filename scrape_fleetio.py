"""
Crawl fleetio.com and export every page's content to fleetio_pages.docx
Each page gets a section with heading, URL, meta description, and body text.
"""
import re, sys, time
from collections import deque
from urllib.parse import urljoin, urlparse

try:
    import requests
except ImportError:
    sys.exit("Missing dependency. Run:  pip install requests")

try:
    from bs4 import BeautifulSoup
except ImportError:
    sys.exit("Missing dependency. Run:  pip install beautifulsoup4")

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
except ImportError:
    sys.exit("Missing dependency. Run:  pip install python-docx")

DOMAIN = "fleetio.com"
START  = "https://www.fleetio.com/"
OUT    = "fleetio_pages.docx"
DELAY  = 0.5
MAX    = 500

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (X11; Linux x86_64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/122.0 Safari/537.36"
    )
}

SKIP_EXTS = re.compile(
    r"\.(pdf|png|jpg|jpeg|gif|svg|ico|webp|mp4|zip|xml|json|css|js)(\?.*)?$",
    re.IGNORECASE,
)

session = requests.Session()
session.headers.update(HEADERS)


def same_domain(url):
    h = urlparse(url).netloc.lstrip("www.")
    return h == DOMAIN or h.endswith("." + DOMAIN)


def clean(text):
    return " ".join(text.split())


def scrape(url):
    try:
        r = session.get(url, timeout=15, allow_redirects=True)
        r.raise_for_status()
        if "text/html" not in r.headers.get("Content-Type", ""):
            return None, []
    except Exception as e:
        print(f"  SKIP {url}: {e}")
        return None, []

    soup = BeautifulSoup(r.text, "lxml")

    links = []
    for a in soup.find_all("a", href=True):
        href = urljoin(url, a["href"].split("#")[0].strip())
        if same_domain(href) and not SKIP_EXTS.search(href):
            links.append(href)

    for tag in soup(["script", "style", "noscript", "header", "footer", "nav"]):
        tag.decompose()

    title     = clean(soup.title.get_text()) if soup.title else url
    meta_desc = ""
    md = soup.find("meta", attrs={"name": "description"})
    if md:
        meta_desc = clean(md.get("content", ""))
    h1   = clean(soup.h1.get_text()) if soup.h1 else ""
    h2s  = [clean(h.get_text()) for h in soup.find_all("h2")]
    body = clean(soup.get_text(" ", strip=True))

    return {"url": url, "title": title, "meta_desc": meta_desc,
            "h1": h1, "h2s": h2s, "body": body}, links


# ── Build the Word document ──────────────────────────────────────────────────

def main():
    doc = Document()
    doc.core_properties.title = "Fleetio.com — Full Site Content"

    # Cover heading
    doc.add_heading("Fleetio.com — Full Site Content", 0)
    doc.add_paragraph(f"Scraped {time.strftime('%Y-%m-%d')}")
    doc.add_page_break()

    visited = set()
    queue   = deque([START])
    count   = 0

    while queue and len(visited) < MAX:
        url = queue.popleft()
        if url in visited:
            continue
        visited.add(url)

        print(f"[{len(visited):>3}] {url}")
        data, links = scrape(url)

        if data:
            count += 1
            # Page title as Heading 1
            doc.add_heading(data["title"] or url, level=1)

            # URL in grey
            p = doc.add_paragraph()
            run = p.add_run(data["url"])
            run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
            run.font.size = Pt(9)

            # Meta description
            if data["meta_desc"]:
                p2 = doc.add_paragraph()
                p2.add_run("Description: ").bold = True
                p2.add_run(data["meta_desc"])

            # H1 / H2s
            if data["h1"]:
                doc.add_heading(data["h1"], level=2)
            for h2 in data["h2s"]:
                doc.add_heading(h2, level=3)

            # Body text
            doc.add_paragraph(data["body"][:5000])
            doc.add_page_break()

        for link in links:
            if link not in visited:
                queue.append(link)

        time.sleep(DELAY)

    doc.save(OUT)
    print(f"\nDone — {count} pages written to {OUT}")


if __name__ == "__main__":
    main()
